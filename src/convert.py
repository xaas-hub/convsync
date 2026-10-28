#!/usr/bin/env python3
"""
Convert files that LLM tooling cannot read as they are (PDF, XLSX, XLS, ODS, DOCX,
ODT) into Markdown under _md/.

markitdown is not used here: its core dependency `magika` pulls in `onnxruntime`,
whose prebuilt wheels are compiled with AVX2 and abort with SIGILL on CPUs without
it. pdftotext (poppler), openpyxl, xlrd, python-docx and odfdo need no such
runtime.

The tree to convert is $CONVSYNC_ROOT (default /data, the volume of the image).

Usage: convert.py [--force]
"""

import hashlib
import os
import re
import subprocess
import sys
from pathlib import Path

import docx
import odfdo
import xlrd
from docx.table import Table
from openpyxl import load_workbook

ROOT = Path(os.environ.get("CONVSYNC_ROOT", "/data")).resolve()
OUT = ROOT / "_md"
OCR = ROOT / "_ocr"
EXCLUDED_DIRS = {"_md", "_ocr"}

# Provenance stamp written on the first line of every generated file. The digest
# covers the body only, so a hand edit can be told apart from a stale rebuild.
MARKER = "<!-- generated from {source} sha256:{digest} -->"
MARKER_RE = re.compile(r"^<!-- generated from .* sha256:([0-9a-f]{64}) -->\n")


def digest_of(body: str) -> str:
    return hashlib.sha256(body.encode("utf-8")).hexdigest()


def read_generated(dst: Path) -> tuple[str | None, str]:
    """Return the recorded digest (None if unstamped) and the body on disk."""
    text = dst.read_text(encoding="utf-8")
    match = MARKER_RE.match(text)
    if match is None:
        return None, text.strip("\n")
    return match.group(1), text[match.end() :].strip("\n")


def ocr_variant(src: Path) -> Path | None:
    """The copy of a scanned PDF that ocr.sh enriched with a text layer."""
    if src.suffix.lower() != ".pdf":
        return None
    candidate = OCR / src.relative_to(ROOT).with_suffix(".ocr.pdf")
    return candidate if candidate.is_file() else None


def pdf_to_markdown(src: Path) -> str:
    # An image-only scan carries no text layer: read the OCR copy instead, so
    # the result lands at the same path as any other document.
    src = ocr_variant(src) or src
    # pdftotext -layout keeps labels and values on the same visual line;
    # pdfminer groups text boxes by column and breaks form-like PDFs.
    result = subprocess.run(
        ["pdftotext", "-layout", "-enc", "UTF-8", "-nopgbrk", str(src), "-"],
        capture_output=True,
        text=True,
        check=True,
    )
    text = result.stdout.strip()
    # Fenced block: the alignment is significant, don't let Markdown collapse it.
    return f"```\n{text}\n```" if text else ""


def markdown_table(rows: list[list[str]]) -> list[str]:
    """Lines of a Markdown table from escaped cell text; the first row is the header."""
    rows = [row for row in rows if any(row)]
    if not rows:
        return []

    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]

    return [
        "| " + " | ".join(rows[0]) + " |",
        "|" + "---|" * width,
        *("| " + " | ".join(row) + " |" for row in rows[1:]),
    ]


def sheets_to_markdown(sheets) -> str:
    """A "## <name>" section with a table per non-empty sheet, from (name, rows) pairs."""
    chunks = []

    for name, values in sheets:
        rows = [
            ["" if cell is None else str(cell).strip().replace("|", "\\|") for cell in row]
            for row in values
        ]
        table = markdown_table(rows)
        if not table:
            continue

        chunks.append(f"## {name}")
        chunks.append("")
        chunks.extend(table)
        chunks.append("")

    return "\n".join(chunks).strip()


def xlsx_to_markdown(src: Path) -> str:
    # data_only=True returns the values cached by Excel/LibreOffice, not formulas.
    workbook = load_workbook(src, data_only=True, read_only=True)
    text = sheets_to_markdown(
        (sheet.title, sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets
    )
    workbook.close()
    return text


def xls_value(cell: xlrd.sheet.Cell, datemode: int):
    """The value of a legacy XLS cell as openpyxl would return it for an XLSX."""
    # BIFF stores every number, dates included, as a float.
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate.xldate_as_datetime(cell.value, datemode)
    if cell.ctype == xlrd.XL_CELL_NUMBER and cell.value.is_integer():
        return int(cell.value)
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return xlrd.error_text_from_code.get(cell.value)
    return cell.value


def xls_to_markdown(src: Path) -> str:
    # Formula cells hold the result cached by the application that saved the file.
    # xlrd prints its warnings on its logfile: keep them off the summary on stdout.
    sheets = []
    with xlrd.open_workbook(src, logfile=sys.stderr) as workbook:
        for sheet in workbook.sheets():
            rows = [[xls_value(c, workbook.datemode) for c in row] for row in sheet.get_rows()]
            sheets.append((sheet.name, rows))
    return sheets_to_markdown(sheets)


def ods_to_markdown(src: Path) -> str:
    sheets = []
    for table in odfdo.Document(src).body.get_tables():
        # LibreOffice declares every sheet as large as Calc allows, through
        # repeated rows and columns of cells that carry only a style: expanded,
        # that is billions of cells. aggressive=True drops the styled ones too.
        table.rstrip(aggressive=True)
        # The values cached by the application, not the formulas.
        sheets.append((table.name, table.get_values()))
    return sheets_to_markdown(sheets)


def docx_to_markdown(src: Path) -> str:
    # Body only: headers, footers, footnotes and text boxes live in other parts
    # of the package. Headings are told apart by the built-in style names, which
    # Word stores in English whatever the language of its UI; list items by their
    # style or, as LibreOffice writes them, by numbering set on the paragraph.
    blocks = []

    for item in docx.Document(src).iter_inner_content():
        if isinstance(item, Table):
            # A merged cell is repeated in every column it spans.
            rows = [
                [cell.text.strip().replace("|", "\\|").replace("\n", " ") for cell in row.cells]
                for row in item.rows
            ]
            table = markdown_table(rows)
            if table:
                blocks.append("\n".join(table))
            continue

        text = item.text.strip()
        if not text:
            continue

        style = item.style.name if item.style is not None else ""
        level = style.removeprefix("Heading ")
        if style == "Title":
            text = f"# {text}"
        elif level != style and level.isdigit():
            text = "#" * min(int(level), 6) + f" {text}"
        elif style.startswith("List Number"):
            text = f"1. {text}"
        # python-docx has no public API for paragraph numbering: read the element.
        elif style.startswith("List") or item._p.pPr is not None and item._p.pPr.numPr is not None:
            text = f"- {text}"
        blocks.append(text)

    return "\n\n".join(blocks)


def odt_to_markdown(src: Path) -> str:
    # odfdo renders headings, lists, tables and emphasis itself. Embedded images
    # come out as links into the archive (Pictures/...), which do not resolve.
    return odfdo.Document(src).to_markdown().strip()


CONVERTERS = {
    ".docx": docx_to_markdown,
    ".ods": ods_to_markdown,
    ".odt": odt_to_markdown,
    ".pdf": pdf_to_markdown,
    ".xls": xls_to_markdown,
    ".xlsx": xlsx_to_markdown,
}


def remove_orphans() -> int:
    """Delete generated files whose source document no longer exists."""
    removed = 0
    if not OUT.is_dir():
        return removed

    for dst in sorted(OUT.rglob("*.md")) + sorted(OUT.rglob("*.md.new")):
        # "_md/2021/degiro/x.pdf.md" -> "2021/degiro/x.pdf"
        rel = dst.relative_to(OUT)
        if rel.suffix == ".new":
            rel = rel.with_suffix("")
        src = ROOT / rel.with_suffix("")
        if src.exists():
            continue
        dst.unlink()
        print(f"REMOVED: {dst.relative_to(ROOT)}")
        removed += 1

    # Drop directories left empty, deepest first.
    for path in sorted(OUT.rglob("*"), reverse=True):
        if path.is_dir() and not any(path.iterdir()):
            path.rmdir()

    return removed


def main() -> int:
    force = "--force" in sys.argv[1:]
    if not ROOT.is_dir():
        print(f"Not a directory: {ROOT} (mount the documents there)", file=sys.stderr)
        return 2
    removed = remove_orphans()
    converted = skipped = conflicts = failed = 0

    for dirpath, dirnames, filenames in os.walk(ROOT):
        dirnames[:] = [
            d for d in dirnames if d not in EXCLUDED_DIRS and not d.startswith(".")
        ]

        for filename in sorted(filenames):
            # The owner file Word leaves next to an open document: a few bytes
            # that share its extension but are no DOCX.
            if filename.startswith("~$"):
                continue

            src = Path(dirpath) / filename
            convert = CONVERTERS.get(src.suffix.lower())
            if convert is None:
                continue

            rel = src.relative_to(ROOT)
            # Keep the original extension in the name: "File inventario.pdf" and
            # "File inventario.xlsx" coexist in the same folder.
            dst = OUT / f"{rel}.md"

            # A re-run of ocr.sh changes the OCR copy, not the original.
            mtime = src.stat().st_mtime
            ocr = ocr_variant(src)
            if ocr is not None:
                mtime = max(mtime, ocr.stat().st_mtime)

            if not force and dst.exists() and dst.stat().st_mtime >= mtime:
                skipped += 1
                continue

            try:
                text = convert(src)
            except Exception as exc:
                print(f"FAILED: {rel} ({exc})", file=sys.stderr)
                failed += 1
                continue

            # An image-only scan yields nothing until ocr.sh has run.
            if not text:
                print(f"EMPTY (scanned PDF, run convsync ocr?): {rel}", file=sys.stderr)

            conflict = dst.with_name(dst.name + ".new")
            stamp = MARKER.format(
                source=f"{rel} (ocr)" if ocr is not None else rel,
                digest=digest_of(text),
            )

            if dst.exists():
                stored, body = read_generated(dst)
                # An unstamped file predates this check and is assumed generated.
                if stored is not None and stored != digest_of(body) and body != text:
                    conflict.write_text(f"{stamp}\n\n{text}\n", encoding="utf-8")
                    print(f"CONFLICT (edited by hand): {rel}", file=sys.stderr)
                    conflicts += 1
                    continue

            dst.parent.mkdir(parents=True, exist_ok=True)
            dst.write_text(f"{stamp}\n\n{text}\n", encoding="utf-8")
            # The hand edit and the converter agree again, or never diverged.
            if conflict.exists():
                conflict.unlink()
            converted += 1

    print(
        f"converted: {converted} | skipped: {skipped} | removed: {removed} | "
        f"conflicts: {conflicts} | failed: {failed}"
    )
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
